import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1] if len(sys.argv) > 1 else "步驟第二版.pptx"
p = Presentation(path)
print(f"File: {path}")
print(f"Slides: {len(p.slides)}")
print(f"Size: {Emu(p.slide_width).inches:.2f}in x {Emu(p.slide_height).inches:.2f}in")
for i, slide in enumerate(p.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            texts.append(shape.text_frame.text.strip().replace("\n", " / "))
    head = texts[0][:60] if texts else "(no text)"
    print(f"  Slide {i}: {head}")